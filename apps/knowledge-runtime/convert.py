"""Deterministic text-first conversion. No URL fetches, macros or formula evaluation.

OCR is the one exception to "text first", and only for PDF pages that have an image but
no text layer: the page is rasterised by pdftoppm and read by Tesseract (ind+eng), both
local binaries in this container, never a service. A page that has text keeps its text;
an OCR'd page is marked as such in the Markdown and in the mapping, and the conversion
carries a warning that says which pages to proof-read against the original. Tesseract
absent from the image means the old behaviour: a scanned page is refused, not emptied.
"""
from __future__ import annotations
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
from html.parser import HTMLParser
import sys
import zipfile
from pathlib import PurePosixPath
from defusedxml import ElementTree as ET

MAX_SOURCE = 50 * 1024 * 1024
MAX_MARKDOWN = 2 * 1024 * 1024
# OCR is slow and heavy on one CPU: a bound on pages, a bound per page, and text that is
# too thin to be a page (a stamp, a logo) is still refused rather than published.
MAX_OCR_PAGES = 10
OCR_PAGE_SECONDS = 30
OCR_MIN_CHARS = 20
OCR_LANGS = 'ind+eng'

def ocr_available():
    return bool(shutil.which('tesseract') and shutil.which('pdftoppm'))

def ocr_page(source, number):
    """One page -> text, or None when there is nothing readable on it."""
    with tempfile.TemporaryDirectory(prefix='ocr-') as work:
        pdf_path = work + '/in.pdf'
        with open(pdf_path, 'wb') as stream:
            stream.write(source)
        # 200 dpi greyscale: enough for body text, a few MB per A4 page.
        subprocess.run(['pdftoppm', '-f', str(number), '-l', str(number), '-r', '200', '-gray',
                        '-png', '-singlefile', pdf_path, work + '/page'],
                       check=True, timeout=OCR_PAGE_SECONDS, stdout=subprocess.DEVNULL,
                       stderr=subprocess.DEVNULL, env={'PATH': '/usr/bin:/bin', 'LANG': 'C.UTF-8'})
        done = subprocess.run(['tesseract', work + '/page.png', 'stdout', '-l', OCR_LANGS, '--psm', '3'],
                              check=True, timeout=OCR_PAGE_SECONDS, capture_output=True,
                              env={'PATH': '/usr/bin:/bin', 'LANG': 'C.UTF-8', 'OMP_THREAD_LIMIT': '1'})
    text = done.stdout.decode('utf-8', 'replace').replace('\x0c', '').strip()
    return text if len(re.sub(r'\s+', '', text)) >= OCR_MIN_CHARS else None

def page_has_image(page):
    resources = page.get('/Resources', {})
    if hasattr(resources, 'get_object'):
        resources = resources.get_object()
    xobjects = resources.get('/XObject', {}) if resources else {}
    if hasattr(xobjects, 'get_object'):
        xobjects = xobjects.get_object()
    for ref in (xobjects or {}).values():
        if ref.get_object().get('/Subtype') == '/Image':
            return True
    return False
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class ConversionError(Exception):
    def __init__(self, code):
        self.code = code
        super().__init__(code)

class Writer:
    def __init__(self):
        self.parts = []
        self.mappings = []
        self.warnings = []
        self.lines = 1
        self.bytes = 0

    def add(self, text, kind, locator):
        text = text + '\n\n'
        size = len(text.encode('utf-8'))
        if self.bytes + size > MAX_MARKDOWN or len(self.mappings) >= 20000:
            raise ConversionError('complexity')
        self.parts.append(text)
        self.mappings.append({'kind': kind, 'locator': locator[:200], 'markdownStart': self.lines,
                              'markdownEnd': self.lines + text.count('\n') - 1})
        self.lines += text.count('\n')
        self.bytes += size

    def warning(self, code):
        if code not in self.warnings:
            self.warnings.append(code)

    def result(self, source):
        markdown = ''.join(self.parts)
        if not markdown.strip():
            raise ConversionError('empty_text')
        return {'pipeline': 'canonical-v2', 'sourceHash': hashlib.sha256(source).hexdigest(),
                'markdownHash': hashlib.sha256(markdown.encode()).hexdigest(), 'markdown': markdown,
                'mappings': self.mappings, 'warnings': self.warnings}

def escape(value):
    # All source text remains data; no uploaded link, image or inline HTML is executable.
    return re.sub(r'([\\`*_{}\[\]<>()#!|])', r'\\\1', str(value)).replace('\r', '').replace('\n', ' ⏎ ')

def literal(text):
    runs = [len(x) for x in re.findall(r'`+', text)]
    fence = '`' * max(3, max(runs, default=0) + 1)
    return fence + 'text\n' + text + '\n' + fence

def office_archive(source, kind):
    try:
        archive = zipfile.ZipFile(io.BytesIO(source))
        infos = archive.infolist()
        if len(infos) > 5000 or len({i.filename for i in infos}) != len(infos):
            raise ConversionError('unsafe_archive')
        total = 0
        for info in infos:
            name = info.filename
            parts = PurePosixPath(name).parts
            if name.startswith('/') or '\\' in name or ':' in name or '..' in parts or (info.external_attr >> 16) & 0o170000 == 0o120000 or info.flag_bits & 1:
                raise ConversionError('unsafe_archive')
            total += info.file_size
            if total > 100 * 1024 * 1024 or info.file_size > 20 * 1024 * 1024 or info.file_size > max(1000000, info.compress_size * 200):
                raise ConversionError('unsafe_archive')
            if any(x in name.lower() for x in ['vbaproject', '/embeddings/', '/activex/', '/externallinks/']):
                raise ConversionError('active_content')
            if name.endswith(('.xml', '.rels')):
                root = ET.fromstring(archive.read(info))
                if name.endswith('.rels'):
                    for rel in root:
                        if rel.get('TargetMode') == 'External' and not rel.get('Type', '').endswith('/hyperlink'):
                            raise ConversionError('active_content')
        expected = 'word/document.xml' if kind == 'DOCX' else 'xl/workbook.xml'
        if expected not in archive.namelist() or '[Content_Types].xml' not in archive.namelist():
            raise ConversionError('invalid_file')
        return archive
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError('invalid_file') from exc

def paragraph_text(node):
    parts = []
    for item in node.iter():
        if item.tag == W + 't':
            parts.append(item.text or '')
        elif item.tag == W + 'tab':
            parts.append('\t')
        elif item.tag in (W + 'br', W + 'cr'):
            parts.append('\n')
        elif item.tag in (W + 'instrText', W + 'fldChar', W + 'object', W + 'altChunk'):
            raise ConversionError('active_content')
    return ''.join(parts)

def numbering_context(archive):
    styles = {}
    if 'word/styles.xml' in archive.namelist():
        for style in ET.fromstring(archive.read('word/styles.xml')).findall(W + 'style'):
            styles[style.get(W + 'styleId')] = style
    nums, abstracts = {}, {}
    if 'word/numbering.xml' in archive.namelist():
        root = ET.fromstring(archive.read('word/numbering.xml'))
        for abstract in root.findall(W + 'abstractNum'):
            abstracts[abstract.get(W + 'abstractNumId')] = abstract
        for num in root.findall(W + 'num'):
            ref = num.find(W + 'abstractNumId')
            if ref is not None:
                nums[num.get(W + 'numId')] = (abstracts.get(ref.get(W + 'val')), num)
    return styles, nums, {}

def numbering_prefix(node, style_id, context):
    styles, nums, counters = context
    pr = node.find('./' + W + 'pPr/' + W + 'numPr')
    seen = set()
    while pr is None and style_id in styles and style_id not in seen:
        seen.add(style_id)
        style = styles[style_id]
        pr = style.find('./' + W + 'pPr/' + W + 'numPr')
        parent = style.find(W + 'basedOn')
        style_id = None if parent is None else parent.get(W + 'val')
    if pr is None:
        return ''
    num_id_node, level_node = pr.find(W + 'numId'), pr.find(W + 'ilvl')
    if num_id_node is None:
        raise ConversionError('complexity')
    num_id = num_id_node.get(W + 'val')
    if num_id == '0':
        return ''
    level = int(level_node.get(W + 'val', '0')) if level_node is not None else 0
    if num_id not in nums or not 0 <= level <= 8:
        raise ConversionError('complexity')
    abstract, num = nums[num_id]
    if abstract is None:
        raise ConversionError('complexity')
    definition = next((d for d in abstract.findall(W + 'lvl') if d.get(W + 'ilvl') == str(level)), None)
    start_override = None
    for override in num.findall(W + 'lvlOverride'):
        if override.get(W + 'ilvl') == str(level):
            if override.find(W + 'lvl') is not None:
                definition = override.find(W + 'lvl')
            start_override = override.find(W + 'startOverride')
    if definition is None:
        raise ConversionError('complexity')
    fmt, pattern = definition.find(W + 'numFmt'), definition.find(W + 'lvlText')
    if fmt is None or pattern is None:
        raise ConversionError('complexity')
    if fmt.get(W + 'val') == 'bullet':
        return '- '
    if fmt.get(W + 'val') != 'decimal':
        raise ConversionError('complexity')
    start = start_override if start_override is not None else definition.find(W + 'start')
    initial = int(start.get(W + 'val', '1')) if start is not None else 1
    key = (num_id, level)
    counters[key] = counters.get(key, initial - 1) + 1
    for deeper in range(level + 1, 9):
        counters.pop((num_id, deeper), None)
    template = pattern.get(W + 'val', '')
    if not re.fullmatch(r'(?:%[1-9]|[.() \-]){1,40}', template):
        raise ConversionError('complexity')
    def replace(match):
        required = (num_id, int(match.group(1)) - 1)
        if required not in counters:
            raise ConversionError('complexity')
        return str(counters[required])
    return re.sub(r'%([1-9])', replace, template) + ' '

def docx(source, out):
    with office_archive(source, 'DOCX') as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        body = root.find(W + 'body')
        if body is None:
            raise ConversionError('invalid_file')
        numbering = numbering_context(archive)
        para = table = 0
        if root.find('.//' + W + 'drawing') is not None or root.find('.//' + W + 'pict') is not None:
            raise ConversionError('complexity')
        for node in body:
            if node.tag == W + 'p':
                para += 1
                text = paragraph_text(node)
                if not text.strip():
                    continue
                style = node.find('./' + W + 'pPr/' + W + 'pStyle')
                style_id = '' if style is None else style.get(W + 'val', '')
                heading = re.match(r'Heading([1-6])$', style_id, re.I)
                prefix = '#' * int(heading.group(1)) + ' ' if heading else ''
                prefix += numbering_prefix(node, style_id, numbering)
                out.add(prefix + escape(text), 'paragraph', f'Paragraf {para}')
            elif node.tag == W + 'tbl':
                table += 1
                rows = []
                for tr in node.findall(W + 'tr'):
                    cells = []
                    for tc in tr.findall(W + 'tc'):
                        if tc.find('./' + W + 'tcPr/' + W + 'gridSpan') is not None or tc.find('./' + W + 'tcPr/' + W + 'vMerge') is not None:
                            raise ConversionError('complexity')
                        cells.append(' ; '.join(escape(paragraph_text(p)) for p in tc.findall(W + 'p')))
                    rows.append(cells)
                if len(rows) > 2000 or any(len(row) > 50 for row in rows):
                    raise ConversionError('complexity')
                if rows:
                    columns = max(map(len, rows))
                    rows = [row + [''] * (columns - len(row)) for row in rows]
                    lines = ['| ' + ' | '.join(row) + ' |' for row in rows]
                    lines.insert(1, '| ' + ' | '.join(['---'] * columns) + ' |')
                    out.add('\n'.join(lines), 'table', f'Tabel {table}, baris 1–{len(rows)}')
            elif node.tag != W + 'sectPr':
                raise ConversionError('complexity')
        for name in archive.namelist():
            if re.match(r'word/(header\d+|footer\d+|footnotes|endnotes)\.xml$', name):
                other = ET.fromstring(archive.read(name))
                for i, p in enumerate(other.iter(W + 'p'), 1):
                    text = paragraph_text(p)
                    if text.strip():
                        out.add(escape(text), 'paragraph', f'{name.split("/")[-1]} paragraf {i}')
        out.warning('Bandingkan paragraf, tabel dan catatan dengan original; penomoran dasar dipertahankan, sel gabungan/struktur kompleks ditolak.')

def xlsx(source, out):
    with office_archive(source, 'XLSX'):
        pass
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    cached = load_workbook(io.BytesIO(source), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(source), read_only=True, data_only=False, keep_links=False)
    try:
        if len(cached.worksheets) > 20:
            raise ConversionError('complexity')
        found_value = False
        for ws, fs in zip(cached.worksheets, formulas.worksheets):
            if ws.max_row > 2000 or ws.max_column > 50:
                raise ConversionError('complexity')
            out.add('## Sheet: ' + escape(ws.title), 'sheet', f'Sheet {ws.title}')
            for row_no, (row, formula_row) in enumerate(zip(ws.iter_rows(), fs.iter_rows()), 1):
                values = []
                for cell, formula in zip(row, formula_row):
                    if formula.data_type == 'f' and cell.value is None:
                        raise ConversionError('missing_cached_value')
                    value = cell.value
                    if hasattr(value, 'isoformat'):
                        value = value.isoformat()
                    rendered = '' if value is None else escape(value)
                    if isinstance(value, (int, float)) and cell.number_format not in ('General', '0'):
                        rendered += ' [format sumber: ' + escape(cell.number_format) + ']'
                    values.append(rendered)
                if not any(values):
                    continue
                found_value = True
                # Each row keeps its exact cell coordinate; no invented header or formula value.
                text = ' | '.join(f'{get_column_letter(i+1)}{row_no}: {v}' for i, v in enumerate(values))
                out.add(text, 'sheet', f'Sheet {ws.title}, A{row_no}:{get_column_letter(len(values))}{row_no}')
        if not found_value:
            raise ConversionError('empty_text')
        out.warning('Nilai cache asli dipakai; formula tidak dihitung. Periksa unit dan koordinat sel pada original.')
    finally:
        cached.close()
        formulas.close()

def pdf(source, out):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(source), strict=True)
        if reader.is_encrypted:
            raise ConversionError('encrypted')
        root = reader.trailer['/Root']
        if any(key in root for key in ['/OpenAction', '/AA', '/AcroForm']):
            raise ConversionError('active_content')
        names = root.get('/Names', {})
        if hasattr(names, 'get_object'):
            names = names.get_object()
        if any(key in names for key in ['/JavaScript', '/EmbeddedFiles']):
            raise ConversionError('active_content')
        if len(reader.pages) > 200:
            raise ConversionError('complexity')
        found = False
        ocr_pages = []
        for i, page in enumerate(reader.pages, 1):
            for ref in page.get('/Annots', []):
                annotation = ref.get_object()
                if any(key in annotation for key in ['/A', '/AA', '/JS', '/FS']):
                    raise ConversionError('active_content')
            if '/AA' in page:
                raise ConversionError('active_content')
            if '/Contents' not in page:
                raise ConversionError('empty_text')
            text = page.extract_text(extraction_mode='layout')
            if not text or not text.strip():
                # A page with no text layer: OCR when it carries an image and OCR is
                # installed; otherwise refuse -- mixed scanned/text PDFs must not
                # silently lose a page.
                if not (page_has_image(page) and ocr_available()):
                    raise ConversionError('empty_text')
                if len(ocr_pages) >= MAX_OCR_PAGES:
                    raise ConversionError('ocr_limit')
                try:
                    text = ocr_page(source, i)
                except subprocess.TimeoutExpired as exc:
                    raise ConversionError('complexity') from exc
                if not text:
                    raise ConversionError('empty_text')
                ocr_pages.append(i)
                found = True
                out.add(f'## Halaman {i} (OCR)\n\n' + literal(text), 'page-ocr', f'Halaman {i}')
                continue
            found = True
            out.add(f'## Halaman {i}\n\n' + literal(text), 'page', f'Halaman {i}')
        if not found:
            raise ConversionError('empty_text')
        if ocr_pages:
            pages = ', '.join(str(p) for p in ocr_pages)
            out.warning(f'Halaman {pages} dibaca dengan OCR (Tesseract, ind+eng) dari gambar pindaian. Cocokkan angka, nama, dan tanda baca dengan original sebelum mengajukan; OCR bisa salah membaca.')
        if len(ocr_pages) < len(reader.pages):
            out.warning('Ekstraksi teks berlayout untuk halaman berteks. Tinjau urutan kolom dan tabel pada setiap halaman; tidak ada bounding box buatan.')
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError('invalid_file') from exc

class HtmlText(HTMLParser):
    """Text and block structure from HTML, nothing else: no attribute, script, style,
    link target or embedded object survives. Blocks become paragraphs, headings, list
    items and simple tables; everything inside is escaped as literal text."""
    SKIP = {'script', 'style', 'noscript', 'template', 'svg', 'math', 'iframe', 'object', 'embed', 'head'}
    BLOCK = {'p', 'div', 'section', 'article', 'header', 'footer', 'main', 'aside', 'nav', 'blockquote',
             'figure', 'figcaption', 'dd', 'dt', 'summary', 'details', 'address', 'hr'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []      # (kind, prefix, text)
        self.buffer = []
        self.prefix = ''
        self.kind = 'paragraph'
        self.skip = 0
        self.lists = []       # 'ul' | 'ol' with counters
        self.table = None     # rows in progress
        self.row = None
        self.cell = None
        self.pre = 0

    def flush(self):
        text = ''.join(self.buffer)
        text = text if self.pre else re.sub(r'\s+', ' ', text).strip()
        if text.strip():
            self.blocks.append((self.kind, self.prefix, text))
        self.buffer = []
        self.prefix = ''
        self.kind = 'paragraph'

    def handle_starttag(self, tag, attrs):
        if self.skip or tag in self.SKIP:
            if tag in self.SKIP:
                self.skip += 1
            return
        if tag == 'table':
            self.flush()
            self.table = []
        elif tag == 'tr' and self.table is not None:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []
        elif self.table is not None:
            return
        elif re.fullmatch(r'h[1-6]', tag):
            self.flush()
            self.prefix = '#' * int(tag[1]) + ' '
        elif tag in ('ul', 'ol'):
            self.flush()
            self.lists.append([tag, 0])
        elif tag == 'li':
            self.flush()
            if self.lists:
                kind, n = self.lists[-1]
                self.lists[-1][1] = n + 1
                self.prefix = '  ' * (len(self.lists) - 1) + ('- ' if kind == 'ul' else f'{n + 1}. ')
        elif tag == 'pre':
            self.flush()
            self.pre += 1
            self.kind = 'code'
        elif tag == 'br':
            self.buffer.append('\n' if self.pre else ' ⏎ ')
        elif tag in self.BLOCK:
            self.flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in ('td', 'th') and self.cell is not None and self.row is not None:
            self.row.append(re.sub(r'\s+', ' ', ''.join(self.cell)).strip())
            self.cell = None
        elif tag == 'tr' and self.row is not None and self.table is not None:
            if self.row:
                self.table.append(self.row)
            self.row = None
        elif tag == 'table' and self.table is not None:
            rows = self.table
            self.table = None
            if rows:
                self.blocks.append(('table', '', rows))
        elif self.table is not None:
            return
        elif tag in ('ul', 'ol'):
            self.flush()
            if self.lists:
                self.lists.pop()
        elif tag == 'pre':
            self.flush()
            self.pre = max(0, self.pre - 1)
        elif re.fullmatch(r'h[1-6]', tag) or tag == 'li' or tag in self.BLOCK:
            self.flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.cell is not None:
            self.cell.append(data)
        elif self.table is None:
            self.buffer.append(data)


def html(source, out):
    try:
        text = source.decode('utf-8')
    except UnicodeDecodeError:
        raise ConversionError('invalid_file')
    if '\x00' in text:
        raise ConversionError('invalid_file')
    parser = HtmlText()
    parser.feed(text)
    parser.close()
    parser.flush()
    blocks = parser.blocks
    if len(blocks) > 20000:
        raise ConversionError('complexity')
    para = table = 0
    for kind, prefix, content in blocks:
        if kind == 'table':
            table += 1
            rows = content
            if len(rows) > 2000 or any(len(row) > 50 for row in rows):
                raise ConversionError('complexity')
            columns = max(map(len, rows))
            rows = [[escape(c) for c in row] + [''] * (columns - len(row)) for row in rows]
            lines = ['| ' + ' | '.join(row) + ' |' for row in rows]
            lines.insert(1, '| ' + ' | '.join(['---'] * columns) + ' |')
            out.add('\n'.join(lines), 'table', f'Tabel {table}, baris 1–{len(rows)}')
        elif kind == 'code':
            para += 1
            out.add(literal(content.strip('\n')), 'paragraph', f'Blok kode {para}')
        else:
            para += 1
            out.add(prefix + escape(content), 'paragraph', f'Paragraf {para}')
    out.warning('Teks diambil dari HTML; tautan, gambar, skrip dan gaya dibuang; tabel gabungan disederhanakan.')


def convert(source, kind):
    if not source or len(source) > MAX_SOURCE:
        raise ConversionError('complexity')
    out = Writer()
    if kind == 'PDF' and source.startswith(b'%PDF-'):
        pdf(source, out)
    elif kind in ('DOCX', 'XLSX') and source.startswith(b'PK\x03\x04'):
        (docx if kind == 'DOCX' else xlsx)(source, out)
    elif kind == 'HTML' and re.match(rb'^\s*(?:\xef\xbb\xbf)?\s*<', source[:64]):
        html(source, out)
    else:
        raise ConversionError('invalid_file')
    return out.result(source)

if __name__ == '__main__':
    import resource
    # PDF may OCR up to MAX_OCR_PAGES; every other format keeps the tight budget.
    cpu = 20 + (MAX_OCR_PAGES * 12 if sys.argv[1] == 'PDF' else 0)
    resource.setrlimit(resource.RLIMIT_CPU, (cpu, cpu))
    resource.setrlimit(resource.RLIMIT_AS, (512*1024*1024, 512*1024*1024))
    resource.setrlimit(resource.RLIMIT_FSIZE, (6*1024*1024, 6*1024*1024))
    try:
        with open(sys.argv[2], 'rb') as stream:
            result = convert(stream.read(MAX_SOURCE + 1), sys.argv[1])
        print(json.dumps(result, ensure_ascii=False))
    except ConversionError as exc:
        print(json.dumps({'code': exc.code}))
        sys.exit(2)
    except Exception:
        print(json.dumps({'code': 'invalid_file'}))
        sys.exit(2)
