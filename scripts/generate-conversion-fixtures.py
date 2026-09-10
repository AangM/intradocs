"""Deterministic synthetic test fixtures, never organizational data.
Requires test-only reportlab/python-docx in addition to converter dependencies.
"""
from pathlib import Path
import io
import json
import zipfile
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'fixtures/uploads'
OUT.mkdir(exist_ok=True, parents=True)

def pdf(name, columns=False):
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(595,842), invariant=1)
    c.setTitle('IntraDocs synthetic conversion fixture')
    c.setFont('Helvetica-Bold',16)
    c.drawString(44,790,'Panduan sintetis — bukan kebijakan resmi')
    c.setFont('Helvetica',11)
    c.drawString(44,750,'1. Verifikasi koneksi sebelum perubahan.')
    c.drawString(44,730,'2. Cadangkan konfigurasi: retensi 30 hari.')
    c.drawString(44,710,'SLA uji: 99.90%; respons maksimal 15 menit.')
    if columns:
        c.drawString(44,670,'Kolom A: 128 Mbps')
        c.drawString(330,670,'Kolom B: 64 Mbps')
        c.drawString(44,650,'Urutan sumber tetap lokal.')
        c.drawString(330,650,'Bandingkan dengan original.')
    c.showPage()
    c.setFont('Helvetica',11)
    c.drawString(44,790,'Halaman dua: uji hasil, lalu catat audit.')
    c.save()
    (OUT/name).write_bytes(stream.getvalue())
    return stream.getvalue()

source=pdf('panduan-demo.pdf')
pdf('multikolom-demo.pdf',True)
writer=PdfWriter();writer.append(PdfReader(io.BytesIO(source)));writer.encrypt('synthetic-only');writer.write(OUT/'encrypted-demo.pdf')
writer=PdfWriter();writer.add_blank_page(width=595,height=842);writer.write(OUT/'scanned-empty-demo.pdf')
doc=Document();doc.add_heading('Panduan Konversi Sintetis',0)
doc.add_paragraph('Contoh Unicode: jaringan, pemulihan, cafés, 日本語.')
doc.add_heading('Langkah aman',1)
doc.add_paragraph('Verifikasi koneksi.',style='List Number')
doc.add_paragraph('Cadangkan konfigurasi.',style='List Number')
doc.add_paragraph('Uji pemulihan dan catat hasil.',style='List Number')
table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
for cell,text in zip(table.rows[0].cells,['Layanan','SLA','Respons']):cell.text=text
for cell,text in zip(table.add_row().cells,['VPN demo','99,90%','15 menit']):cell.text=text
doc.save(OUT/'panduan-demo.docx')
wb=Workbook();ws=wb.active;ws.title='SLA Demo'
for row in [['Layanan','SLA (%)','Respons (menit)'],['VPN demo',99.9,15],['Backup demo',99.5,30]]:ws.append(row)
for cell in ws[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1D4ED8')
for row in ws.iter_rows(min_row=2):
    for cell in row:cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(vertical='center')
for col,width in [('A',24),('B',16),('C',24)]:ws.column_dimensions[col].width=width
ws.freeze_panes='A2';ws['B2'].number_format=ws['B3'].number_format='0.00';ws['C2'].number_format=ws['C3'].number_format='0'
t=Table(displayName='SyntheticSLA',ref='A1:C3');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
wb.save(OUT/'sla-demo.xlsx')
# An intentionally uncached formula is a rejection fixture, not an example of a completed workbook.
ws['C4']='=C2+C3';wb.save(OUT/'formula-no-cache.xlsx')
with zipfile.ZipFile(OUT/'hostile-office.docx','w') as z:
    z.writestr('../outside.xml','<x/>');z.writestr('[Content_Types].xml','<Types/>');z.writestr('word/document.xml','<x/>')
(OUT/'corrupt-demo.pdf').write_bytes(b'%PDF-1.7\nnot a valid document')
print('Generated synthetic conversion and explicit rejection fixtures.')
